"""
Export an Excel workbook with per-module tool / DECA statistics.

Sheets:
  1. Modules        — one row per module: total tools, unique, shared + who they're shared with
  2. PNs par module — one row per PN×module: DECA count, other modules that also use it
  3. Matrice        — module × module heat-map of shared PNs

Usage:
    python -m scripts.export_module_stats [--out stats_modules.xlsx]
"""
import argparse
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent.parent))
from db.db import fetchall

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
C_HEADER   = "1F3864"   # dark navy
C_HEADER_T = "FFFFFF"
C_UNIQUE   = "D9EAD3"   # light green
C_SHARED   = "FCE5CD"   # light orange
C_HEAT_LO  = "FFFFFF"
C_HEAT_HI  = "E06C75"   # red for high sharing
C_ALT      = "F3F3F3"
C_ACCENT   = "2A82DA"

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _header_style(ws, row: int, values: list, widths: list | None = None):
    for col, val in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(bold=True, color=C_HEADER_T, size=10)
        c.fill = _fill(C_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def _cell(ws, row, col, value, fill_hex=None, bold=False, align="left", wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = BORDER
    if fill_hex:
        c.fill = _fill(fill_hex)
    return c


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def load_data():
    rows = fetchall("""
        SELECT
            t.marquage,
            t.pn_short,
            t.modules,
            t.complexity_flag,
            t.exclusion_reason
        FROM tools t
        WHERE t.pn_short IS NOT NULL
          AND t.modules IS NOT NULL
          AND t.modules != ''
          AND t.exclusion_reason IS NULL
    """)

    # pn_short → set of modules
    pn_modules: dict[str, set] = defaultdict(set)
    # module → list of (marquage, pn_short)
    mod_tools: dict[str, list] = defaultdict(list)
    # pn_short → list of marquages
    pn_decas: dict[str, list] = defaultdict(list)

    for r in rows:
        mods = [m.strip() for m in (r["modules"] or "").split(",") if m.strip()]
        pn = r["pn_short"]
        marquage = r["marquage"]
        for mod in mods:
            pn_modules[pn].add(mod)
            mod_tools[mod].append((marquage, pn))
        pn_decas[pn].append(marquage)

    return pn_modules, mod_tools, pn_decas


# ---------------------------------------------------------------------------
# Sheet 1 — Module summary
# ---------------------------------------------------------------------------

def sheet_modules(wb, pn_modules, mod_tools, pn_decas):
    ws = wb.create_sheet("Modules")
    ws.freeze_panes = "A2"

    headers = [
        "Module", "Total outils (DECAs)",
        "PNs distincts", "PNs uniques au module",
        "PNs partagés", "Partagé avec (modules)",
        "DECAs total",
    ]
    widths = [12, 22, 18, 22, 16, 45, 14]
    _header_style(ws, 1, headers, widths)

    modules = sorted(mod_tools.keys())

    for r_idx, mod in enumerate(modules, 2):
        tools = mod_tools[mod]
        pns_in_mod = set(pn for _, pn in tools)
        deca_count = len(tools)

        unique_pns = {pn for pn in pns_in_mod if pn_modules[pn] == {mod}}
        shared_pns = pns_in_mod - unique_pns

        # Who else shares those PNs?
        shared_with: dict[str, int] = defaultdict(int)
        for pn in shared_pns:
            for other_mod in pn_modules[pn]:
                if other_mod != mod:
                    shared_with[other_mod] += 1
        shared_str = "  |  ".join(
            f"{m} ({n} PN{'s' if n>1 else ''})"
            for m, n in sorted(shared_with.items())
        ) if shared_with else "—"

        alt = C_ALT if r_idx % 2 == 0 else None

        _cell(ws, r_idx, 1, mod,              bold=True, align="center")
        _cell(ws, r_idx, 2, deca_count,       align="center", fill_hex=alt)
        _cell(ws, r_idx, 3, len(pns_in_mod),  align="center", fill_hex=alt)
        _cell(ws, r_idx, 4, len(unique_pns),  align="center",
              fill_hex=C_UNIQUE if unique_pns else alt)
        _cell(ws, r_idx, 5, len(shared_pns),  align="center",
              fill_hex=C_SHARED if shared_pns else alt)
        _cell(ws, r_idx, 6, shared_str,       fill_hex=alt, wrap=True)
        _cell(ws, r_idx, 7, sum(len(pn_decas[pn]) for pn in pns_in_mod),
              align="center", fill_hex=alt)
        ws.row_dimensions[r_idx].height = 18

    # Totals row
    last = len(modules) + 2
    total_tools = sum(len(v) for v in mod_tools.values())
    ws.cell(last, 1, "TOTAL").font = Font(bold=True, size=10)
    ws.cell(last, 2, total_tools).font = Font(bold=True, size=10)
    ws.cell(last, 2).alignment = Alignment(horizontal="center")
    for c in range(1, 8):
        ws.cell(last, c).fill = _fill("D9D9D9")
        ws.cell(last, c).border = BORDER

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


# ---------------------------------------------------------------------------
# Sheet 2 — PNs par module
# ---------------------------------------------------------------------------

def sheet_pns(wb, pn_modules, mod_tools, pn_decas):
    ws = wb.create_sheet("PNs par module")
    ws.freeze_panes = "A2"

    headers = ["Module", "PN", "DECAs", "Autres modules", "Complexité"]
    widths  = [12, 16, 10, 45, 16]
    _header_style(ws, 1, headers, widths)

    r_idx = 2
    for mod in sorted(mod_tools.keys()):
        pns_in_mod = sorted({pn for _, pn in mod_tools[mod]})
        for pn in pns_in_mod:
            other_mods = sorted(pn_modules[pn] - {mod})
            n_decas = len(pn_decas[pn])
            if len(pn_modules[pn]) > 1:
                complexity = "partagé"
                fill = C_SHARED
            elif n_decas > 1:
                complexity = "multi_deca"
                fill = "D9EAD3"
            else:
                complexity = "unique"
                fill = None

            alt = C_ALT if r_idx % 2 == 0 else None
            _cell(ws, r_idx, 1, mod,  bold=True, align="center")
            _cell(ws, r_idx, 2, pn,   align="center")
            _cell(ws, r_idx, 3, n_decas, align="center", fill_hex=fill or alt)
            _cell(ws, r_idx, 4, ", ".join(other_mods) if other_mods else "—",
                  fill_hex=alt, wrap=True)
            _cell(ws, r_idx, 5, complexity, align="center",
                  fill_hex=fill or alt)
            ws.row_dimensions[r_idx].height = 18
            r_idx += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


# ---------------------------------------------------------------------------
# Sheet 3 — Matrice de partage
# ---------------------------------------------------------------------------

def sheet_matrix(wb, pn_modules, mod_tools):
    ws = wb.create_sheet("Matrice partage")
    modules = sorted(mod_tools.keys())
    n = len(modules)

    # Build sharing matrix
    matrix: dict[tuple, int] = defaultdict(int)
    for pn, mods in pn_modules.items():
        mods_l = sorted(mods)
        for i in range(len(mods_l)):
            for j in range(i + 1, len(mods_l)):
                matrix[(mods_l[i], mods_l[j])] += 1

    max_val = max(matrix.values()) if matrix else 1

    def heat_color(val: int) -> str:
        if val == 0:
            return C_HEAT_LO
        ratio = val / max_val
        r = int(0xE0 * ratio + 0xFF * (1 - ratio))
        g = int(0x6C * ratio + 0xFF * (1 - ratio))
        b = int(0x75 * ratio + 0xFF * (1 - ratio))
        return f"{r:02X}{g:02X}{b:02X}"

    # Header row
    ws.cell(1, 1, "Module ↓  /  →").font = Font(bold=True, size=9)
    ws.cell(1, 1).fill = _fill(C_HEADER)
    ws.cell(1, 1).font = Font(bold=True, color=C_HEADER_T, size=9)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(1, 1).border = BORDER
    ws.column_dimensions["A"].width = 14

    for j, mod in enumerate(modules, 2):
        c = ws.cell(1, j, mod)
        c.font = Font(bold=True, color=C_HEADER_T, size=9)
        c.fill = _fill(C_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = 10
    ws.row_dimensions[1].height = 28

    for i, mod_row in enumerate(modules, 2):
        c = ws.cell(i, 1, mod_row)
        c.font = Font(bold=True, color=C_HEADER_T, size=9)
        c.fill = _fill(C_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        ws.row_dimensions[i].height = 20

        for j, mod_col in enumerate(modules, 2):
            if mod_row == mod_col:
                cell = ws.cell(i, j, "—")
                cell.fill = _fill("D9D9D9")
            else:
                key = tuple(sorted([mod_row, mod_col]))
                val = matrix.get(key, 0)
                cell = ws.cell(i, j, val if val else "")
                if val:
                    cell.fill = _fill(heat_color(val))
                    cell.font = Font(bold=val > max_val * 0.5, size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

    # Legend
    legend_row = n + 3
    ws.cell(legend_row, 1, "Nombre de PNs partagés entre deux modules").font = Font(italic=True, size=9)
    ws.cell(legend_row + 1, 1, f"Max = {max_val} PNs partagés").font = Font(italic=True, size=9)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="stats_modules.xlsx")
    args = parser.parse_args()

    print("Chargement des données…")
    pn_modules, mod_tools, pn_decas = load_data()

    if not mod_tools:
        print("Aucun outil avec module trouvé dans la DB. Avez-vous rechargé les sources ?")
        sys.exit(1)

    print(f"  {len(mod_tools)} modules, {len(pn_modules)} PNs distincts")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    print("Génération feuille Modules…")
    sheet_modules(wb, pn_modules, mod_tools, pn_decas)

    print("Génération feuille PNs par module…")
    sheet_pns(wb, pn_modules, mod_tools, pn_decas)

    print("Génération matrice de partage…")
    sheet_matrix(wb, pn_modules, mod_tools)

    out_path = Path(args.out)
    wb.save(out_path)
    print(f"\n✓ Fichier généré : {out_path.resolve()}")


if __name__ == "__main__":
    main()
